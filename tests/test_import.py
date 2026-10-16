import base64
import datetime
import io
import json

import openpyxl

from odoo.tests import tagged
from odoo.tests.common import HttpCase


@tagged('post_install', '-at_install')
class TestChurchImport(HttpCase):
    """Imports Excel : le modèle de l'église, puis un import tout ou rien."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, tracking_disable=True))
        # Les numéros saisis « 07 … » sont ivoiriens : la société doit l'être.
        cls.env.company.country_id = cls.env.ref('base.ci')
        Church = cls.env['equipment.church']
        cls.church = Church.create({'name': 'Église A', 'followup_enabled': True})
        cls.other = Church.create({'name': 'Église B', 'followup_enabled': True})
        cls.manager = cls.env['church.mobile.user'].create({
            'name': 'Resp', 'phone': '+225 07 00 00 09 09', 'role': 'manager', 'church_id': cls.church.id})
        cls.env['church.district'].create({'name': 'Riviera', 'church_id': cls.church.id})
        cls.env['church.district'].create({'name': 'Quartier de B', 'church_id': cls.other.id})
        cls.Import = cls.env['church.import']

    # ─── Outils ───────────────────────────────────────────────────────

    def _template(self, kind, church=None):
        return openpyxl.load_workbook(io.BytesIO(self.Import._template(kind, church or self.church)))

    def _filled(self, kind, rows, church=None):
        """Le modèle de l'église, rempli de ces lignes (dans l'ordre des colonnes données)."""
        book = self._template(kind, church)
        sheet = book.worksheets[0]
        headers = [c.value for c in sheet[1]]
        for r, row in enumerate(rows, start=2):
            for header, value in row.items():
                col = next(i for i, h in enumerate(headers, start=1) if h.rstrip(' *') == header)
                sheet.cell(row=r, column=col, value=value)
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    def _run(self, kind, rows, church=None, user=None):
        return self.Import._run(kind, church or self.church, self._filled(kind, rows, church),
                                'test.xlsx', mobile_user=user)

    def _rpc(self, path, account, **params):
        payload = {'jsonrpc': '2.0', 'method': 'call', 'params': params}
        headers = {'Content-Type': 'application/json',
                   'Authorization': 'Bearer %s' % account._open_session('test')}
        return self.url_open(path, data=json.dumps(payload).encode(), headers=headers).json()['result']

    # ─── Le modèle ────────────────────────────────────────────────────

    def test_template_carries_this_church_lists_only(self):
        book = self._template('prayer_cell')
        self.assertEqual(book.sheetnames, ['Cellules', "Mode d'emploi", 'Listes'])
        headers = [c.value for c in book['Cellules'][1]]
        self.assertIn('Nom de la cellule *', headers, "Les colonnes obligatoires sont marquées")
        lists = [c.value for row in book['Listes'].iter_rows() for c in row if c.value]
        self.assertIn('Riviera', lists)
        self.assertIn('Vendredi', lists)
        self.assertNotIn('Quartier de B', lists, "Le modèle ne montre que l'église de celui qui importe")
        validations = book['Cellules'].data_validations.dataValidation
        self.assertTrue(any('Listes' in (v.formula1 or '') for v in validations), "Listes déroulantes")

    # ─── Tout ou rien ─────────────────────────────────────────────────

    def test_districts_import_and_reimport_creates_no_duplicate(self):
        rows = [{'Nom du quartier': 'Angré', 'Ville': 'Cocody'}, {'Nom du quartier': 'RIVIÉRA'}]
        result = self._run('district', rows)
        self.assertTrue(result['ok'], result)
        self.assertEqual((result['created'], result['updated'], result['unchanged']), (1, 0, 1),
                         "« RIVIÉRA » est le quartier « Riviera » déjà connu")
        again = self._run('district', rows)
        self.assertEqual((again['created'], again['unchanged']), (0, 2))
        self.assertEqual(self.env['church.district'].search_count([('church_id', '=', self.church.id)]), 2)

    def test_cells_are_imported_with_their_leader_account(self):
        result = self._run('prayer_cell', [{
            'Nom de la cellule': 'Cellule Angré',
            'Quartiers couverts': 'Angré Château; riviera',
            'Jour de réunion': 'jeudi',
            'Heure': '19h30',
            'Responsable': 'Fatou Cissé',
            # Excel a lu un nombre : le zéro initial est perdu, l'import le rend.
            'Téléphone du responsable': 700050505,
        }])
        self.assertTrue(result['ok'], result)
        cell = self.env['church.prayer.cell'].search([('name', '=', 'Cellule Angré')])
        self.assertEqual(cell.meeting_day, 'thursday')
        self.assertEqual(cell.meeting_time, '19:30')
        self.assertEqual(sorted(cell.district_ids.mapped('name')), ['Angré Château', 'Riviera'])
        self.assertIn('Angré Château', result['warnings'][0], "Le quartier créé est signalé")
        self.assertEqual(cell.mobile_user_id.role, 'cell_leader')
        self.assertEqual(cell.mobile_user_id.phone_e164, '+2250700050505')
        self.assertEqual(cell.church_id, self.church)

    def test_one_bad_line_imports_nothing_and_says_where(self):
        rows = [
            {'Nom de la cellule': 'C1', 'Jour de réunion': 'Lundi'},
            {'Nom de la cellule': 'C2', 'Jour de réunion': 'Lundi'},
            {'Nom de la cellule': 'C3', 'Jour de réunion': 'Funday', 'Heure': '25:00'},
            {'Nom de la cellule': '', 'Lieu de réunion': 'Chez Paul'},
            {'Nom de la cellule': 'c1'},
        ]
        result = self._run('prayer_cell', rows)
        self.assertFalse(result['ok'])
        self.assertEqual(self.env['church.prayer.cell'].search_count([('church_id', '=', self.church.id)]), 0,
                         "Rien n'est écrit")
        text = '\n'.join(result['errors'])
        self.assertIn('Ligne 4 — Jour de réunion', text)
        self.assertIn('Ligne 4 — Heure', text)
        self.assertIn('Ligne 5 — Nom de la cellule : obligatoire', text)
        self.assertIn('Ligne 6 : même fiche que la ligne 2', text)
        log = self.env['church.import.log'].browse(result['log_id'])
        self.assertEqual(log.state, 'refused', "Le refus est gardé dans l'historique, fichier compris")
        self.assertTrue(log.file)

    def test_failure_while_writing_rolls_back_every_line(self):
        """Le numéro du responsable a déjà un compte dans une autre église : la
        ligne échoue à l'écriture, et les lignes déjà écrites sont défaites."""
        self.env['church.mobile.user'].create({
            'name': 'Ailleurs', 'phone': '+225 07 00 00 07 07', 'role': 'evangelist', 'church_id': self.other.id})
        result = self._run('prayer_cell', [
            {'Nom de la cellule': 'C1'},
            {'Nom de la cellule': 'C2', 'Responsable': 'Ailleurs', 'Téléphone du responsable': '07 00 00 07 07'},
        ])
        self.assertFalse(result['ok'])
        self.assertIn('Ligne 3', result['errors'][0])
        self.assertIn('Église B', result['errors'][0], "Le message dit où est l'autre compte")
        self.assertFalse(self.env['church.prayer.cell'].search([('name', 'in', ('C1', 'C2'))]))

    def test_existing_cell_is_completed_never_overwritten(self):
        cell = self.env['church.prayer.cell'].create({
            'name': 'Cellule Riviera', 'church_id': self.church.id, 'leader_name': 'Paul'})
        result = self._run('prayer_cell', [{
            'Nom de la cellule': 'cellule riviera', 'Lieu de réunion': 'Chez Paul', 'Responsable': 'Marc',
            'Créer son compte': 'Non'}])
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['updated'], 1)
        self.assertEqual(cell.meeting_location, 'Chez Paul', "La case vide est remplie")
        self.assertEqual(cell.leader_name, 'Paul', "Rien n'est écrasé")
        self.assertTrue(any('« Paul » conservé' in w for w in result['warnings']))

    def test_age_groups_check_their_choices(self):
        bad = self._run('age_group', [{'Nom du groupe': 'Jeunes', 'Type': 'Jeunes', 'Sexe': 'Tous'}])
        self.assertFalse(bad['ok'])
        self.assertIn('Ligne 2 — Sexe', bad['errors'][0])
        good = self._run('age_group', [{'Nom du groupe': 'Mariés', 'Type': 'mariés', 'Sexe': 'Mixte',
                                        "Tranche d'âge": 'Mariés 25-35 ans', 'Créer son compte': 'Non'}])
        self.assertTrue(good['ok'], good)
        group = self.env['church.age.group'].search([('name', '=', 'Mariés'), ('church_id', '=', self.church.id)])
        self.assertEqual((group.group_type, group.gender), ('married', 'mixed'))
        self.assertTrue(group.age_range_id)

    def test_unreadable_or_foreign_files_are_refused_plainly(self):
        result = self.Import._run('district', self.church, b'pas un classeur', 'liste.xlsx')
        self.assertFalse(result['ok'])
        self.assertIn('pas un classeur Excel', result['errors'][0])
        book = openpyxl.Workbook()
        book.active.append(['Prénom', 'Âge'])
        book.active.append(['Awa', 30])
        buffer = io.BytesIO()
        book.save(buffer)
        result = self.Import._run('district', self.church, buffer.getvalue(), 'autre.xlsx')
        self.assertIn('Aucun en-tête reconnu', result['errors'][0])

    def test_csv_is_accepted(self):
        content = 'Nom du quartier;Ville\nYopougon;Abidjan\n'.encode('cp1252')
        result = self.Import._run('district', self.church, content, 'quartiers.csv')
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['created'], 1)

    def test_dates_and_times_typed_by_excel(self):
        result = self._run('prayer_cell', [{'Nom de la cellule': 'C', 'Heure': datetime.time(18, 45)}])
        self.assertTrue(result['ok'], result)
        self.assertEqual(self.env['church.prayer.cell'].search([('name', '=', 'C')]).meeting_time, '18:45')

    # ─── Depuis l'application ─────────────────────────────────────────

    def test_api_template_and_import_stay_in_the_user_church(self):
        res = self._rpc('/api/church/import/template', self.manager, kind='district', church_id=self.other.id)
        self.assertEqual(res['status'], 'success', res)
        self.assertIn('Église A', res['filename'], "Un responsable reçoit le modèle de SON église")
        content = self._filled('district', [{'Nom du quartier': 'Abobo'}])
        res = self._rpc('/api/church/import/run', self.manager, kind='district', church_id=self.other.id,
                        filename='q.xlsx', content_base64=base64.b64encode(content).decode())
        self.assertEqual(res['status'], 'success', res)
        district = self.env['church.district'].search([('name', '=', 'Abobo')])
        self.assertEqual(district.church_id, self.church)
        self.assertEqual(self.env['church.import.log'].search([], limit=1).mobile_user_id, self.manager)

    def test_api_refusal_lists_the_errors(self):
        content = self._filled('age_group', [{'Nom du groupe': 'X'}])
        res = self._rpc('/api/church/import/run', self.manager, kind='age_group', filename='g.xlsx',
                        content_base64=base64.b64encode(content).decode())
        self.assertEqual(res['code'], 'import_refused')
        self.assertTrue(any('Type : obligatoire' in e for e in res['errors']))
        bad = self._rpc('/api/church/import/run', self.manager, kind='age_group', content_base64='%%%')
        self.assertEqual(bad['code'], 'invalid')

    def test_only_pastor_manager_and_admin_import(self):
        evangelist = self.env['church.evangelist'].create({
            'name': 'Ev', 'phone': '+225 07 00 00 01 01', 'church_id': self.church.id}).mobile_user_id
        res = self._rpc('/api/church/import/kinds', evangelist)
        self.assertEqual(res['code'], 'forbidden')
        admin = self.env['church.mobile.user'].create({
            'name': 'Admin', 'phone': '+225 07 00 00 77 77', 'role': 'super_admin'})
        res = self._rpc('/api/church/import/template', admin, kind='district')
        self.assertEqual(res['code'], 'invalid', "L'administrateur multi-église désigne l'église")
        res = self._rpc('/api/church/import/template', admin, kind='district', church_id=self.other.id)
        self.assertIn('Église B', res['filename'])
