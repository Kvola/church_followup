"""Imports Excel du suivi : un modèle généré pour l'église, un import TOUT OU RIEN.

Le parcours, depuis l'application comme depuis le back-office :

1. télécharger le modèle de SON église : un onglet de saisie, un mode
   d'emploi, et des listes déroulantes remplies avec les cellules, groupes et
   quartiers de cette église (on choisit au lieu de taper) ;
2. le remplir (Excel, WPS, Google Sheets…) ;
3. l'importer. Le fichier est vérifié EN ENTIER avant la moindre écriture. À
   la première erreur rien n'est écrit, et la réponse dit, ligne par ligne,
   pourquoi. Sans erreur, tout est écrit d'un seul tenant.

Une fiche qui existe déjà n'est pas une erreur : elle est COMPLÉTÉE (seuls ses
champs vides sont remplis), jamais écrasée. Ce qui diffère est signalé, sans
bloquer. Réimporter le même fichier ne crée donc aucun doublon.

Chaque tentative, réussie ou refusée, est gardée dans l'historique avec le
fichier d'origine.
"""
import base64
import datetime
import io
import logging
import re
import unicodedata

from psycopg2 import IntegrityError

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

from .church_member import name_tokens
from .church_tools import MAX_PROPOSED_WEEKS

_logger = logging.getLogger(__name__)

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
MAX_SHOWN_ERRORS = 50
GUIDE_SHEET = "Mode d'emploi"
LISTS_SHEET = 'Listes'
LIST_SEPARATOR = ';'
XLSX_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ImportRefused(Exception):
    """Le fichier ne peut pas être importé ; `errors` dit pourquoi."""

    def __init__(self, errors):
        super().__init__('\n'.join(errors))
        self.errors = errors


class _Rollback(Exception):
    """Annule l'écriture en cours : une ligne au moins a échoué."""


def norm_key(value):
    """« Cellule  Angré 8ᵉ » et « cellule angre 8e » donnent la même clé."""
    text = unicodedata.normalize('NFKD', str(value or ''))
    text = ''.join(c for c in text if not unicodedata.combining(c)).lower()
    return ' '.join(re.sub(r'[^0-9a-z]+', ' ', text).split())


# ─── Colonnes ─────────────────────────────────────────────────────────────


class Column:
    """Une colonne du modèle : son en-tête, sa nature, ses valeurs permises.

    `kind` : text, phone, time, choice, bool, list (plusieurs valeurs séparées
    par « ; »), date, int. `choices(env, church)` rend [(libellé, valeur)] ;
    les libellés remplissent la liste déroulante et sont reconnus à l'import
    (la valeur technique aussi, sans accents ni majuscules).
    """

    def __init__(self, key, header, kind='text', required=False, help='', choices=None,
                 example='', width=22, default=None, minimum=None, maximum=None):
        self.key = key
        self.header = header
        self.kind = kind
        self.required = required
        self.help = help
        self.choices = choices
        self.example = example
        self.width = width
        self.default = default
        self.minimum = minimum
        self.maximum = maximum

    @property
    def title(self):
        return '%s *' % self.header if self.required else self.header


def _selection(model, field):
    def choices(env, church):
        return [(label, value) for value, label in env[model]._fields[field]._description_selection(env)]
    return choices


YES_NO = lambda env, church: [('Oui', True), ('Non', False)]  # noqa: E731


# ─── Les imports ──────────────────────────────────────────────────────────


class ImportSpec:
    """Un import : ses colonnes, comment il reconnaît une fiche, comment il écrit."""

    kind = None
    label = None
    sheet = None
    model = None
    intro = ''
    columns = []

    # Reconnaissance d'une fiche existante (dans l'église) : par défaut, le nom.
    def identity(self, values):
        return norm_key(values.get('name'))

    def prepare_context(self, env, church):
        """Ce que l'import charge une fois avant d'écrire (fiches existantes…)."""
        records = env[self.model].with_context(active_test=False).search([('church_id', '=', church.id)])
        return {'existing': {norm_key(r.name): r for r in records}}

    def match(self, env, church, values, ctx):
        """La fiche existante qui correspond à la ligne, ou None. ValueError si ambigu."""
        return ctx['existing'].get(self.identity(values))

    def remember(self, ctx, values, record):
        ctx['existing'][self.identity(values)] = record

    def to_vals(self, env, church, values, row, ctx):
        """Les valeurs du modèle Odoo pour cette ligne (déjà contrôlée)."""
        return {k: v for k, v in values.items() if k in env[self.model]._fields}

    def create_vals(self, values):
        """Ce qu'on ne pose qu'à la création (jamais pour compléter une fiche)."""
        return {}

    def after_write(self, env, church, record, values, created, ctx, row):
        """Ce qui suit l'écriture (compte du responsable, suivi…). Peut lever."""

    def notes(self, ctx):
        """Ce qu'il faut dire en plus du décompte (quartiers créés, suivis démarrés…)."""
        created = ctx.get('created_districts') or []
        return [_("Quartiers créés : %s.") % ', '.join(created)] if created else []


class LeaderColumnsMixin:
    """Responsable d'une cellule ou d'un groupe : nom, téléphone, compte."""

    leader_columns = [
        Column('leader_name', 'Responsable', example='Fatou Cissé', width=24,
               help="Nom et prénom du responsable."),
        Column('leader_phone', 'Téléphone du responsable', kind='phone', example='07 00 05 05 05',
               help="C'est par ce numéro qu'il se connectera à l'application."),
        Column('create_account', 'Créer son compte', kind='bool', choices=YES_NO, default=True,
               example='Oui', width=16,
               help="Oui (par défaut) : son compte de l'application est ouvert, il n'a plus qu'à "
                    "se connecter avec son numéro. Un numéro n'ouvre qu'une église."),
    ]

    def after_write(self, env, church, record, values, created, ctx, row):
        if values.get('create_account', True) and record.leader_phone and not record.mobile_user_id.active:
            record.action_create_leader_account()


def _districts_context(env, church):
    districts = env['church.district'].with_context(active_test=False).search([('church_id', '=', church.id)])
    return {'districts': {norm_key(d.name): d for d in districts}, 'created_districts': []}


def _district(env, church, ctx, name):
    """Le quartier de ce nom dans l'église ; créé s'il n'existe pas (et signalé)."""
    district = ctx['districts'].get(norm_key(name))
    if not district:
        district = env['church.district'].create({'name': name, 'church_id': church.id})
        ctx['districts'][norm_key(name)] = district
        ctx['created_districts'].append(name)
    return district


def _church_names(model, label_extra=None):
    """Les fiches actives de l'église, par nom ; deux homonymes se distinguent."""
    def choices(env, church):
        records = env[model].search([('church_id', '=', church.id)], order='name')
        counts = {}
        for r in records:
            counts[norm_key(r.name)] = counts.get(norm_key(r.name), 0) + 1
        return [('%s (%s)' % (r.name, label_extra(r)) if label_extra and counts[norm_key(r.name)] > 1
                 else r.name, r.id) for r in records]
    return choices


class DistrictImport(ImportSpec):
    kind = 'district'
    label = 'Quartiers'
    sheet = 'Quartiers'
    model = 'church.district'
    intro = "Les quartiers où vivent les membres. Les cellules de prière s'y rattachent."
    columns = [
        Column('name', 'Nom du quartier', required=True, example='Angré 8e tranche', width=30),
        Column('city', 'Ville', example='Cocody'),
        Column('description', 'Description', example='Autour du carrefour Djibi', width=40),
    ]


class PrayerCellImport(LeaderColumnsMixin, ImportSpec):
    kind = 'prayer_cell'
    label = 'Cellules de prière'
    sheet = 'Cellules'
    model = 'church.prayer.cell'
    intro = ("Les cellules de prière de l'église. Un quartier inconnu est créé ; plusieurs "
             "quartiers se séparent par « ; ».")

    @property
    def columns(self):
        return [
            Column('name', 'Nom de la cellule', required=True, example='Cellule Angré 8e tranche', width=30),
            Column('district_names', 'Quartiers couverts', kind='list', example='Angré 8e tranche; Angré Château',
                   width=34, choices=lambda env, church: [
                       (d.name, d.name) for d in env['church.district'].search([('church_id', '=', church.id)])],
                   help="Un ou plusieurs quartiers, séparés par « ; ». Un quartier absent de la liste est créé."),
            Column('meeting_day', 'Jour de réunion', kind='choice', choices=_selection('church.prayer.cell', 'meeting_day'),
                   example='Vendredi', width=16),
            Column('meeting_time', 'Heure', kind='time', example='19:00', width=10,
                   help="Au format 19:00 (ou 19h, 19h30)."),
            Column('meeting_location', 'Lieu de réunion', example='Chez la famille Cissé', width=28),
        ] + self.leader_columns + [
            Column('notes', 'Notes', width=30),
        ]

    def to_vals(self, env, church, values, row, ctx):
        vals = super().to_vals(env, church, values, row, ctx)
        names = values.get('district_names') or []
        if names:
            vals['district_ids'] = [_district(env, church, ctx, name).id for name in names]
        return vals

    def prepare_context(self, env, church):
        return dict(super().prepare_context(env, church), **_districts_context(env, church))


class AgeGroupImport(LeaderColumnsMixin, ImportSpec):
    kind = 'age_group'
    label = "Groupes d'âge"
    sheet = 'Groupes'
    model = 'church.age.group'
    intro = "Les groupes d'âge de l'église (mariés, jeunes, lycéens…)."

    @property
    def columns(self):
        return [
            Column('name', 'Nom du groupe', required=True, example='Mariés 25-35', width=26),
            Column('group_type', 'Type', kind='choice', required=True,
                   choices=_selection('church.age.group', 'group_type'), example='Mariés', width=16),
            Column('gender', 'Sexe', kind='choice', required=True,
                   choices=_selection('church.age.group', 'gender'), example='Mixte', width=12),
            Column('age_range_id', "Tranche d'âge", kind='choice', width=22, example='Mariés 25-35 ans',
                   choices=lambda env, church: [(r.name, r.id) for r in env['church.age.range'].search(
                       ['|', ('church_id', '=', False), ('church_id', '=', church.id)])],
                   help="Facultatif ; surtout pour les groupes de mariés."),
        ] + self.leader_columns + [
            Column('notes', 'Notes', width=30),
        ]


def _phone_identity(values):
    return ('phone', values['phone']) if values.get('phone') else None


class EvangelistImport(ImportSpec):
    kind = 'evangelist'
    label = 'Évangélistes'
    sheet = 'Evangelistes'
    model = 'church.evangelist'
    intro = ("Les évangélistes de l'église. Leur compte de l'application est ouvert : ils se "
             "connectent avec leur numéro. Un numéro n'ouvre qu'une église.")
    columns = [
        Column('name', 'Nom et prénoms', required=True, example='Awa Koné', width=28),
        Column('phone', 'Téléphone', kind='phone', required=True, example='07 00 01 01 01',
               help="Son numéro de connexion à l'application."),
        Column('email', 'E-mail', example='awa.kone@exemple.ci', width=28),
    ]

    def identity(self, values):
        return _phone_identity(values) or ('name', ' '.join(sorted(name_tokens(values.get('name')))))

    def prepare_context(self, env, church):
        Otp = env['otp.code']
        by_phone, by_name = {}, {}
        for ev in env[self.model].with_context(active_test=False).search([('church_id', '=', church.id)]):
            try:
                if ev.phone:
                    by_phone[Otp._norm_phone(ev.phone)] = ev
            except ValidationError:
                pass
            by_name.setdefault(' '.join(sorted(name_tokens(ev.name))), ev)
        return {'by_phone': by_phone, 'by_name': by_name}

    def match(self, env, church, values, ctx):
        found = ctx['by_phone'].get(values.get('phone'))
        if not found:
            other = ctx['by_name'].get(' '.join(sorted(name_tokens(values.get('name')))))
            # Même nom mais autre numéro : un homonyme, pas la même personne.
            if other and not other.phone:
                found = other
        return found

    def remember(self, ctx, values, record):
        ctx['by_phone'][values['phone']] = record
        ctx['by_name'][' '.join(sorted(name_tokens(values.get('name'))))] = record


class PersonImport(ImportSpec):
    """Ce qui est commun aux membres et aux personnes à suivre : une personne."""

    model = 'church.member'

    def person_columns(self, extra_before=(), extra_after=()):
        return list(extra_before) + [
            Column('name', 'Nom', required=True, example='KONAN', width=18),
            Column('first_name', 'Prénoms', required=True, example='Aya Marie', width=22),
            Column('gender', 'Sexe', kind='choice', choices=_selection('church.member', 'gender'),
                   example='Femme', width=10, help="Facultatif : la fiche sera alors « à compléter »."),
            Column('phone', 'Téléphone', kind='phone', example='07 08 09 10 11',
                   help="Sert aussi à reconnaître une personne déjà enregistrée."),
            Column('email', 'E-mail', width=24),
            Column('date_of_birth', 'Date de naissance', kind='date', example='25/12/1995', width=14,
                   help="Au format JJ/MM/AAAA. Distingue deux homonymes."),
            Column('marital_status', 'Situation matrimoniale', kind='choice',
                   choices=_selection('church.member', 'marital_status'), example='Célibataire', width=18),
            Column('profession', 'Profession', example='Commerçante'),
            Column('district_name', 'Quartier', kind='choice', example='Angré 8e tranche', width=24,
                   choices=lambda env, church: [(d.name, d.name) for d in env['church.district'].search(
                       [('church_id', '=', church.id)], order='name')],
                   help="Un quartier absent de la liste est créé."),
            Column('address', "Lieu d'habitation", example='Près de la pharmacie', width=26),
            Column('city', 'Ville', example='Abidjan', width=14),
            Column('salvation_date', 'Date de conversion', kind='date', example='12/05/2026', width=14),
        ] + list(extra_after) + [
            Column('notes', 'Notes', width=30),
        ]

    def identity(self, values):
        return _phone_identity(values) or (
            'name', ' '.join(sorted(name_tokens(values.get('name'), values.get('first_name')))),
            values.get('date_of_birth'))

    def prepare_context(self, env, church):
        pool = env['church.member'].with_context(active_test=False).search([('church_id', '=', church.id)])
        return dict(_districts_context(env, church), pool=pool, started=0, to_assign=0, already=[])

    def match(self, env, church, values, ctx):
        same, _reason, ambiguous = env['church.member']._find_same_person(
            church, values.get('name'), values.get('first_name'), values.get('phone'),
            values.get('date_of_birth'), pool=ctx['pool'])
        if ambiguous:
            raise ValueError(_("plusieurs personnes s'appellent « %s » : indiquez le téléphone ou la "
                               "date de naissance pour savoir laquelle.") % same.display_name_custom)
        return same or None

    def remember(self, ctx, values, record):
        ctx['pool'] |= record

    def to_vals(self, env, church, values, row, ctx):
        vals = super().to_vals(env, church, values, row, ctx)
        if values.get('district_name'):
            vals['district_id'] = _district(env, church, ctx, values['district_name']).id
        return vals


class MemberImport(PersonImport):
    kind = 'member'
    label = 'Membres'
    sheet = 'Membres'
    intro = ("Les membres de l'église (anciens membres), avec leur cellule et leur groupe. Une "
             "personne déjà enregistrée (même numéro, ou même nom et date de naissance) est complétée.")

    @property
    def columns(self):
        return self.person_columns(extra_after=[
            Column('prayer_cell_id', 'Cellule de prière', kind='choice', width=26,
                   choices=_church_names('church.prayer.cell', lambda r: r.meeting_location or r.id),
                   example='Cellule Angré 8e tranche'),
            Column('age_group_id', "Groupe d'âge", kind='choice', width=20,
                   choices=_church_names('church.age.group', lambda r: r.id), example='Mariés 25-35'),
            Column('baptism_date', 'Date de baptême', kind='date', example='04/08/2019', width=14),
            Column('integration_date', "Date d'intégration", kind='date', width=14),
        ])

    def create_vals(self, values):
        return {'member_type': 'old_member', 'person_profile': 'old_member'}


class ToFollowImport(PersonImport):
    kind = 'to_follow'
    label = 'Personnes à suivre'
    sheet = 'A suivre'
    intro = ("Les personnes à suivre : nouvelles personnes ou anciens membres repris en suivi. "
             "Avec un évangéliste, le suivi démarre ; sans, la personne attend dans « À attribuer ».")

    @property
    def columns(self):
        return self.person_columns(
            extra_before=[
                Column('person_profile', 'Profil', kind='choice', default='new', width=18,
                       choices=_selection('church.member', 'person_profile'), example='Nouvelle personne',
                       help="Nouvelle personne (par défaut) ou ancien membre repris en suivi."),
            ],
            extra_after=[
                Column('evangelist_id', 'Évangéliste', kind='choice', width=24,
                       choices=_church_names('church.evangelist', lambda r: r.phone or r.id),
                       example='Awa Koné',
                       help="Facultatif. Indiqué : le suivi démarre. Vide : « À attribuer »."),
                Column('duration_weeks', 'Durée du suivi (semaines)', kind='int', minimum=1,
                       maximum=MAX_PROPOSED_WEEKS, width=14, example='6',
                       help="Facultatif : la durée de l'église s'applique (1 à %d)." % MAX_PROPOSED_WEEKS),
                Column('start_date', 'Début du suivi', kind='date', width=14, example='19/10/2026',
                       help="Facultatif : aujourd'hui."),
            ])

    def to_vals(self, env, church, values, row, ctx):
        vals = super().to_vals(env, church, values, row, ctx)
        # Le profil, le statut et le référent se décident à la création ou au suivi,
        # jamais en « complétant » une fiche.
        for key in ('person_profile', 'evangelist_id', 'duration_weeks', 'start_date'):
            vals.pop(key, None)
        return vals

    def create_vals(self, values):
        return {'person_profile': values.get('person_profile') or 'new', 'member_type': 'new'}

    def after_write(self, env, church, record, values, created, ctx, row):
        member = record
        if member.member_type == 'in_followup' or member.followup_ids.filtered(
                lambda f: f.state in ('in_progress', 'extended')):
            ctx['already'].append(member.display_name_custom)
            return
        evangelist = env['church.evangelist'].browse(values['evangelist_id']) if values.get('evangelist_id') else None
        if evangelist:
            vals = {'church_id': church.id, 'evangelist_id': evangelist.id, 'member_id': member.id,
                    'start_date': values.get('start_date') or fields.Date.context_today(member)}
            if values.get('duration_weeks'):
                vals['duration_weeks'] = values['duration_weeks']
            env['church.followup'].create(vals)
            ctx['started'] += 1
            return
        # Sans évangéliste : la personne attend qu'on lui en confie un.
        changes = {}
        if member.member_type not in ('new', 'interrupted'):  # un suivi interrompu y est déjà
            changes['member_type'] = 'new'
        if values.get('duration_weeks'):
            changes['planned_followup_weeks'] = values['duration_weeks']
        if changes:
            member.write(changes)
        ctx['to_assign'] += 1

    def notes(self, ctx):
        notes = super().notes(ctx)
        notes.append(_("%(s)s suivi(s) démarré(s), %(a)s personne(s) dans « À attribuer ».",
                       s=ctx.get('started', 0), a=ctx.get('to_assign', 0)))
        if ctx.get('already'):
            notes.append(_("Déjà en suivi, laissé(e)s tel quel : %s.") % ', '.join(ctx['already']))
        return notes


SPECS = {spec.kind: spec for spec in (
    DistrictImport(), PrayerCellImport(), AgeGroupImport(),
    EvangelistImport(), MemberImport(), ToFollowImport(),
)}
IMPORT_KINDS = [(spec.kind, spec.label) for spec in SPECS.values()]


# ─── Le moteur ────────────────────────────────────────────────────────────


class ChurchImport(models.AbstractModel):
    _name = 'church.import'
    _description = "Imports Excel du suivi"

    @api.model
    def _spec(self, kind):
        spec = SPECS.get(kind)
        if not spec:
            raise UserError(_("Import inconnu : %s") % kind)
        return spec

    @api.model
    def _kinds_payload(self):
        return [{'kind': s.kind, 'label': s.label, 'description': s.intro} for s in SPECS.values()]

    # ── Le modèle ────────────────────────────────────────────────────────

    @api.model
    def _template_filename(self, kind, church):
        spec = self._spec(kind)
        return 'Modele import %s - %s.xlsx' % (spec.label, church.name)

    @api.model
    def _template(self, kind, church):
        """Le modèle Excel de cet import pour cette église (octets)."""
        import xlsxwriter  # présent dans les images Docker (Dockerfile, Dockerfile-prod)
        from xlsxwriter.utility import xl_col_to_name

        spec = self._spec(kind)
        env = self.sudo().env
        columns = spec.columns
        buffer = io.BytesIO()
        book = xlsxwriter.Workbook(buffer, {'in_memory': True})
        head = book.add_format({'bold': True, 'bg_color': '#1F4E79', 'font_color': 'white',
                                'text_wrap': True, 'valign': 'vcenter', 'border': 1})
        head_required = book.add_format({'bold': True, 'bg_color': '#C55A11', 'font_color': 'white',
                                         'text_wrap': True, 'valign': 'vcenter', 'border': 1})
        as_text = book.add_format({'num_format': '@'})
        title = book.add_format({'bold': True, 'font_size': 14})
        bold = book.add_format({'bold': True})
        wrap = book.add_format({'text_wrap': True, 'valign': 'top'})

        sheet = book.add_worksheet(spec.sheet)
        guide = book.add_worksheet(GUIDE_SHEET)
        lists = book.add_worksheet(LISTS_SHEET)

        # Onglet de saisie : les en-têtes, et c'est tout (l'exemple vit dans le mode d'emploi).
        sheet.set_row(0, 32)
        sheet.freeze_panes(1, 0)
        list_col = 0
        for index, column in enumerate(columns):
            sheet.write(0, index, column.title, head_required if column.required else head)
            text_like = column.kind in ('text', 'phone', 'time', 'list')
            sheet.set_column(index, index, column.width, as_text if text_like else None)
            if column.help:
                sheet.write_comment(0, index, column.help, {'x_scale': 2, 'y_scale': 1.2})
            if column.choices and column.kind in ('choice', 'bool', 'list'):
                labels = [label for label, _value in column.choices(env, church)]
                lists.write(0, list_col, column.header, bold)
                for row, label in enumerate(labels, start=1):
                    lists.write_string(row, list_col, str(label))
                lists.set_column(list_col, list_col, max(16, column.width))
                if labels and column.kind != 'list':  # une liste multiple ne se choisit pas en un clic
                    col_letter = xl_col_to_name(list_col)
                    sheet.data_validation(1, index, MAX_ROWS, index, {
                        'validate': 'list',
                        'source': "='%s'!$%s$2:$%s$%d" % (LISTS_SHEET, col_letter, col_letter, len(labels) + 1),
                        'error_title': column.header,
                        'error_message': "Choisissez une valeur de la liste.",
                    })
                list_col += 1
        if not list_col:
            lists.write(0, 0, "Aucune liste pour cet import.")

        # Mode d'emploi.
        guide.set_column(0, 0, 28)
        guide.set_column(1, 1, 14)
        guide.set_column(2, 2, 30)
        guide.set_column(3, 3, 70)
        guide.write(0, 0, "Import : %s — %s" % (spec.label, church.name), title)
        lines = [
            spec.intro,
            "Remplissez l'onglet « %s », une ligne par fiche, sans toucher aux en-têtes." % spec.sheet,
            "Les colonnes en orange (marquées *) sont obligatoires.",
            "Le fichier est vérifié en entier : à la moindre erreur, RIEN n'est importé et "
            "l'application dit quelle ligne corriger.",
            "Une fiche qui existe déjà (même nom) est complétée : ses cases vides sont remplies, "
            "rien n'est écrasé. Réimporter le fichier ne crée pas de doublon.",
            "Au plus %d lignes par fichier." % MAX_ROWS,
        ]
        for row, text in enumerate(lines, start=2):
            guide.write(row, 0, text)
        row = len(lines) + 3
        for col, text in enumerate(('Colonne', 'Obligatoire', 'Exemple', 'Explication')):
            guide.write(row, col, text, head)
        for column in columns:
            row += 1
            allowed = ''
            if column.choices and column.kind in ('choice', 'bool'):
                labels = [label for label, _value in column.choices(env, church)]
                allowed = ' Valeurs : %s.' % ', '.join(map(str, labels[:30])) if labels else ''
            guide.write(row, 0, column.header, bold)
            guide.write(row, 1, 'Oui' if column.required else '')
            guide.write_string(row, 2, str(column.example or ''))
            guide.write(row, 3, (column.help + allowed).strip(), wrap)
        book.close()
        return buffer.getvalue()

    # ── La lecture du fichier ────────────────────────────────────────────

    @api.model
    def _read_rows(self, spec, content, filename):
        """Les lignes du fichier : [(n° de ligne Excel, {clé: valeur brute})]."""
        if not content:
            raise ImportRefused([_("Le fichier est vide.")])
        if len(content) > MAX_BYTES:
            raise ImportRefused([_("Le fichier dépasse %d Mo.") % (MAX_BYTES // (1024 * 1024))])
        name = (filename or '').lower()
        if name.endswith('.csv'):
            table = self._read_csv(content)
        elif name.endswith('.xls'):
            raise ImportRefused([_("Le format .xls (ancien Excel) n'est pas lu : enregistrez le "
                                   "fichier au format .xlsx.")])
        else:
            table = self._read_xlsx(spec, content)

        by_header = {norm_key(c.header): c for c in spec.columns}
        header_row, mapping = None, {}
        for index, cells in table[:10]:
            found = {}
            for position, cell in enumerate(cells):
                column = by_header.get(norm_key(str(cell or '').replace('*', '')))
                if column:
                    found[position] = column.key
            if found:
                header_row, mapping = index, found
                break
        if header_row is None:
            raise ImportRefused([_("Aucun en-tête reconnu : utilisez le modèle « %s » sans "
                                   "modifier sa première ligne.") % spec.label])
        missing = [c.header for c in spec.columns if c.required and c.key not in mapping.values()]
        if missing:
            raise ImportRefused([_("Colonne obligatoire absente : %s. Utilisez le modèle sans "
                                   "modifier sa première ligne.") % ', '.join(missing)])
        rows = []
        for index, cells in table:
            if index <= header_row:
                continue
            values = {key: cells[position] if position < len(cells) else None
                      for position, key in mapping.items()}
            if all(v is None or str(v).strip() == '' for v in values.values()):
                continue
            rows.append((index, values))
        if not rows:
            raise ImportRefused([_("Aucune ligne à importer : l'onglet « %s » est vide.") % spec.sheet])
        if len(rows) > MAX_ROWS:
            raise ImportRefused([_("%(n)s lignes : c'est plus que les %(max)s permises par fichier. "
                                   "Découpez-le.", n=len(rows), max=MAX_ROWS)])
        return rows

    @staticmethod
    def _read_xlsx(spec, content):
        import openpyxl
        try:
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - fichier corrompu, mot de passe, autre format
            raise ImportRefused([_("Le fichier n'est pas un classeur Excel (.xlsx) lisible.")])
        try:
            names = book.sheetnames
            wanted = next((n for n in names if norm_key(n) == norm_key(spec.sheet)), None)
            if not wanted:
                wanted = next((n for n in names if norm_key(n) not in (norm_key(GUIDE_SHEET), norm_key(LISTS_SHEET))),
                              names[0])
            sheet = book[wanted]
            return [(index, list(cells)) for index, cells in
                    enumerate(sheet.iter_rows(values_only=True), start=1)]
        finally:
            book.close()

    @staticmethod
    def _read_csv(content):
        import csv
        for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=';,\t')
        except csv.Error:  # une seule colonne : rien à deviner
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        return [(index, row) for index, row in enumerate(reader, start=1)]

    # ── Le contrôle des valeurs ──────────────────────────────────────────

    def _convert(self, column, raw, church, choice_maps):
        """La valeur propre d'une case, ou ValueError avec le message à montrer."""
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if column.required:
                raise ValueError(_("obligatoire."))
            return column.default
        kind = column.kind
        if kind == 'text':
            if isinstance(raw, float) and raw.is_integer():
                raw = int(raw)
            return ' '.join(str(raw).split())
        if kind == 'phone':
            return self._convert_phone(raw)
        if kind == 'time':
            return self._convert_time(raw)
        if kind == 'list':
            parts = re.split(r'[;,\n]', str(raw))
            seen, result = set(), []
            for part in (' '.join(p.split()) for p in parts):
                if part and norm_key(part) not in seen:
                    seen.add(norm_key(part))
                    result.append(part)
            return result
        if kind in ('choice', 'bool'):
            key = norm_key(raw)
            mapping = choice_maps[column.key]
            if key in mapping:
                return mapping[key]
            if kind == 'bool' and key in ('oui', 'o', 'yes', 'y', 'x', '1', 'vrai', 'true'):
                return True
            if kind == 'bool' and key in ('non', 'n', 'no', '0', 'faux', 'false'):
                return False
            labels = [label for label in choice_maps[column.key + ':labels']]
            raise ValueError(_("« %(v)s » n'est pas une valeur permise (%(allowed)s).",
                               v=raw, allowed=', '.join(map(str, labels[:12])) or _('liste vide')))
        if kind == 'date':
            return self._convert_date(raw)
        if kind == 'int':
            try:
                number = float(str(raw).replace(',', '.'))
            except ValueError:
                raise ValueError(_("« %s » n'est pas un nombre.") % raw)
            if not number.is_integer():
                raise ValueError(_("« %s » n'est pas un nombre entier.") % raw)
            number = int(number)
            if (column.minimum is not None and number < column.minimum) or \
                    (column.maximum is not None and number > column.maximum):
                raise ValueError(_("%(n)s : attendu entre %(lo)s et %(hi)s.",
                                   n=number, lo=column.minimum, hi=column.maximum))
            return number
        return raw

    def _convert_phone(self, raw):
        if isinstance(raw, (int, float)):
            # Excel a lu un nombre : il a mangé le zéro initial (07… → 7…).
            digits = str(int(raw))
            raw = '0' + digits if len(digits) == 9 else digits
        text = str(raw).strip()
        digits = re.sub(r'\D', '', text)
        if len(digits) < 8:
            raise ValueError(_("« %s » n'est pas un numéro complet.") % text)
        try:
            return self.env['otp.code'].sudo()._norm_phone(text)
        except ValidationError:
            raise ValueError(_("« %s » n'est pas un numéro de téléphone lisible.") % text)

    @staticmethod
    def _convert_time(raw):
        if isinstance(raw, datetime.datetime):
            raw = raw.time()
        if isinstance(raw, datetime.time):
            return '%02d:%02d' % (raw.hour, raw.minute)
        if isinstance(raw, (int, float)) and 0 <= raw < 1:  # fraction de jour d'Excel
            minutes = round(raw * 24 * 60)
            return '%02d:%02d' % (minutes // 60 % 24, minutes % 60)
        match = re.fullmatch(r'\s*(\d{1,2})\s*(?:[:hH.]\s*(\d{1,2})?)?\s*', str(raw))
        if match:
            hour, minute = int(match.group(1)), int(match.group(2) or 0)
            if hour < 24 and minute < 60:
                return '%02d:%02d' % (hour, minute)
        raise ValueError(_("« %s » n'est pas une heure (ex. 19:00).") % raw)

    @staticmethod
    def _convert_date(raw):
        if isinstance(raw, datetime.datetime):
            return raw.date()
        if isinstance(raw, datetime.date):
            return raw
        text = str(raw).strip()
        for pattern in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y'):
            try:
                return datetime.datetime.strptime(text, pattern).date()
            except ValueError:
                continue
        raise ValueError(_("« %s » n'est pas une date (ex. 25/12/1990).") % text)

    # ── L'import ─────────────────────────────────────────────────────────

    @api.model
    def _run(self, kind, church, content, filename, mobile_user=None):
        """Importe le fichier, tout ou rien.

        Rend {ok, created, updated, unchanged, errors, warnings, message, log_id}.
        Ne lève pas pour un fichier mauvais : le refus est une réponse.
        """
        spec = self._spec(kind)
        env = self.sudo().with_context(tracking_disable=True, mail_create_nolog=True,
                                       mail_notrack=True).env
        stats = {'created': 0, 'updated': 0, 'unchanged': 0}
        errors, warnings, ctx = [], [], {}
        try:
            rows = self._read_rows(spec, content, filename)
            prepared = self._check_rows(spec, env, church, rows, errors)
            if errors:
                raise ImportRefused(errors)
            ctx = spec.prepare_context(env, church)
            ctx['warnings'] = warnings
            try:
                with env.cr.savepoint():
                    self._write_rows(spec, env, church, prepared, ctx, stats, errors, warnings)
                    if errors:
                        raise _Rollback()
            except _Rollback:
                env.invalidate_all()
                raise ImportRefused(errors)
        except ImportRefused as refused:
            errors = refused.errors
            stats = {'created': 0, 'updated': 0, 'unchanged': 0}
            warnings = []

        ok = not errors
        if ok:
            warnings[:0] = spec.notes(ctx)
            message = _("%(label)s : %(c)s créé(s), %(u)s complété(s), %(n)s déjà à jour.",
                        label=spec.label, c=stats['created'], u=stats['updated'], n=stats['unchanged'])
        else:
            message = _("Rien n'a été importé : %s erreur(s) à corriger dans le fichier.") % len(errors)
        shown = errors[:MAX_SHOWN_ERRORS]
        if len(errors) > MAX_SHOWN_ERRORS:
            shown.append(_("… et %s autre(s).") % (len(errors) - MAX_SHOWN_ERRORS))
        log = self.env['church.import.log'].sudo().create({
            'church_id': church.id,
            'kind': kind,
            'state': 'done' if ok else 'refused',
            'mobile_user_id': mobile_user.id if mobile_user else False,
            'user_id': self.env.user.id if not mobile_user else False,
            'filename': filename or 'import.xlsx',
            'file': base64.b64encode(content or b''),
            'created_count': stats['created'],
            'updated_count': stats['updated'],
            'unchanged_count': stats['unchanged'],
            'error_count': len(errors),
            'message': message,
            'errors': '\n'.join(errors) or False,
            'warnings': '\n'.join(warnings) or False,
        })
        return dict(stats, ok=ok, errors=shown, warnings=warnings[:MAX_SHOWN_ERRORS],
                    message=message, log_id=log.id)

    def _check_rows(self, spec, env, church, rows, errors):
        """Contrôle chaque case de chaque ligne, sans rien écrire."""
        columns = spec.columns
        choice_maps = {}
        for column in columns:
            if column.choices and column.kind in ('choice', 'bool'):
                pairs = column.choices(env, church)
                mapping = {}
                for label, value in pairs:
                    mapping[norm_key(label)] = value
                    if isinstance(value, str):
                        mapping.setdefault(norm_key(value), value)
                choice_maps[column.key] = mapping
                choice_maps[column.key + ':labels'] = [label for label, _value in pairs]
        prepared, seen = [], {}
        for row, raw in rows:
            values, row_ok = {}, True
            for column in columns:
                try:
                    values[column.key] = self._convert(column, raw.get(column.key), church, choice_maps)
                except ValueError as e:
                    errors.append(_("Ligne %(row)s — %(col)s : %(msg)s", row=row, col=column.header, msg=e))
                    row_ok = False
            if not row_ok:
                continue
            identity = spec.identity(values)
            if identity in seen:
                errors.append(_("Ligne %(row)s : même fiche que la ligne %(first)s (« %(name)s »).",
                                row=row, first=seen[identity], name=values.get('name') or identity))
                continue
            seen[identity] = row
            prepared.append((row, values))
        return prepared

    def _write_rows(self, spec, env, church, prepared, ctx, stats, errors, warnings):
        """Écrit les lignes contrôlées ; une ligne qui échoue est notée, pas fatale ici."""
        Model = env[spec.model]
        for row, values in prepared:
            try:
                with env.cr.savepoint():
                    record = spec.match(env, church, values, ctx)
                    vals = spec.to_vals(env, church, values, row, ctx)
                    if record:
                        changes = self._complete(record, vals, row, warnings)
                        if changes:
                            record.write(changes)
                            stats['updated'] += 1
                        else:
                            stats['unchanged'] += 1
                        created = False
                    else:
                        vals.update(spec.create_vals(values), church_id=church.id)
                        vals = {k: v for k, v in vals.items() if v not in (None, '', [])}
                        record = Model.create(vals)
                        stats['created'] += 1
                        created = True
                    spec.remember(ctx, values, record)
                    spec.after_write(env, church, record, values, created, ctx, row)
            except ValueError as e:  # rapprochement impossible (homonymes…)
                errors.append(_("Ligne %(row)s : %(msg)s", row=row, msg=e))
            except (UserError, ValidationError) as e:
                errors.append(_("Ligne %(row)s : %(msg)s", row=row, msg=e.args[0] if e.args else e))
            except IntegrityError:
                errors.append(_("Ligne %s : cette fiche existe déjà sous une autre forme.") % row)

    @staticmethod
    def _complete(record, vals, row, warnings):  # noqa: C901
        """Les champs à remplir sur une fiche existante : ses cases vides seulement."""
        changes = {}
        if 'active' in record._fields and not record.active:
            changes['active'] = True
        for name, value in vals.items():
            if value in (None, '', []) or name not in record._fields:
                continue
            field = record._fields[name]
            current = record[name]
            if field.type == 'many2many':
                missing = [i for i in value if i not in current.ids]
                if missing:
                    changes[name] = [(4, i) for i in missing]
                continue
            current_value = current.id if field.type == 'many2one' else current
            if not current_value:
                changes[name] = value
            elif field.type in ('char', 'text') and norm_key(current_value) == norm_key(value):
                continue  # même texte, à la casse ou aux accents près
            elif name == 'phone' and _same_phone(record.env, current_value, value):
                continue  # même numéro, écrit autrement
            elif current_value != value:
                shown = current.display_name if field.type == 'many2one' else current
                warnings.append(_("Ligne %(row)s — %(col)s : « %(old)s » conservé (le fichier dit « %(new)s »).",
                                  row=row, col=field.string, old=shown, new=value if field.type != 'many2one'
                                  else record.env[field.comodel_name].browse(value).display_name))
        return changes


def _same_phone(env, a, b):
    Otp = env['otp.code'].sudo()
    try:
        return Otp._norm_phone(a) == Otp._norm_phone(b)
    except ValidationError:
        return False


class ChurchImportLog(models.Model):
    _name = 'church.import.log'
    _description = "Historique des imports du suivi"
    _order = 'create_date desc, id desc'
    _rec_name = 'filename'

    church_id = fields.Many2one('equipment.church', string='Église', required=True, index=True,
                                ondelete='cascade')
    kind = fields.Selection(IMPORT_KINDS, string='Import', required=True)
    state = fields.Selection([('done', 'Importé'), ('refused', 'Refusé')], string='Résultat', required=True)
    mobile_user_id = fields.Many2one('church.mobile.user', string="Par (application)", ondelete='set null')
    user_id = fields.Many2one('res.users', string='Par (back-office)', ondelete='set null')
    filename = fields.Char(string='Fichier')
    file = fields.Binary(string='Fichier importé', attachment=True)
    created_count = fields.Integer(string='Créés')
    updated_count = fields.Integer(string='Complétés')
    unchanged_count = fields.Integer(string='Déjà à jour')
    error_count = fields.Integer(string='Erreurs')
    message = fields.Char(string='Résumé')
    errors = fields.Text(string='Erreurs détaillées')
    warnings = fields.Text(string='Remarques')
